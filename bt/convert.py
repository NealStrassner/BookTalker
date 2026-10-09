"""Any readable file -> a PDF, so every format gets the same page view, layout pass,
voice and highlighting. Results are cached per file.

    PDF                          as is
    EPUB MOBI FB2 XPS CBZ        PyMuPDF lays them out as book pages
    PNG JPG TIFF (scans)         one page per image; read with OCR
    AZW3 AZW MOBI PRC PDB        our Kindle reader (bt/mobi.py); copy-protected books are refused
    CHM                          our CHM reader with its own LZX decompressor (bt/chm.py)
    DJVU                         its text layer, each word where it was printed (bt/djvu.py)
    WPD WP                       our WordPerfect reader when Word isn't there
    DOCX DOC RTF ODT             Word's exact look when Word is installed — otherwise our own
    PPTX PPT ODP                 readers (bt/readers.py), so nothing depends on Office
    ODS EML MHT TXTZ HTMLZ FBZ   our readers
    TXT MD TEX HTML …            typeset as book pages; headings become chapters
"""
import html
import os
import re
import subprocess
import tempfile

import pymupdf

from .layout import EXACT_WORDS, cache_dir, doc_id, stext_words

NATIVE = {".epub", ".mobi", ".fb2", ".xps", ".oxps", ".svg",
          ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}   # images: OCR'd
WORD = {".docx", ".doc", ".rtf", ".odt", ".docm", ".dotx", ".dot", ".ott", ".wpd", ".wp", ".wp4", ".wp5", ".wp6"}
SLIDES = {".pptx", ".ppt", ".odp", ".pptm", ".ppsx", ".pps", ".potx", ".otp"}
TEXT = {".txt", ".md", ".markdown", ".text", ".log", ".csv", ".tsv", ".rst", ".adoc", ".asciidoc",
        ".org", ".nfo", ".tex"}
WEB = {".html", ".htm", ".xhtml"}
KINDLE = {".azw3", ".azw", ".kf8", ".prc", ".pdb"}
CODE = {".py", ".pyw", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".c", ".h", ".cpp", ".cc", ".cxx", ".hpp", ".hh",
        ".cs", ".java", ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".kts", ".lua", ".sh", ".bash", ".zsh", ".ps1",
        ".psm1", ".bat", ".cmd", ".vbs", ".sql", ".r", ".pl", ".pm", ".scala", ".vb", ".bas", ".pas", ".dart", ".m",
        ".mm", ".f90", ".jl", ".hs", ".ex", ".exs", ".clj", ".erl", ".ml", ".fs", ".nim", ".zig", ".groovy", ".ino",
        ".css", ".scss", ".asm", ".s", ".mq4", ".mq5", ".pine", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg",
        ".xml", ".gradle", ".cmake", ".mk", ".dockerfile"}   # source files: read as code (bt/code.py)
OTHER = {".chm", ".ods", ".ots", ".eml", ".mht", ".mhtml", ".txtz", ".htmlz", ".fbz", ".lit", ".djvu", ".djv",
         ".cbr", ".cbz", ".cbt"}
PICTURES = (".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".jp2", ".jxr")
SUPPORTED = {".pdf"} | NATIVE | WORD | SLIDES | TEXT | WEB | KINDLE | OTHER | CODE
OPEN_FILTER = "Readable files (" + " ".join("*" + e for e in sorted(SUPPORTED)) + ");;PDF (*.pdf);;All files (*.*)"

BOOK_PAGE = pymupdf.Rect(0, 0, 432, 648)      # 6 x 9 in, a trade-book page
COMIC_MARK = "BookTalker comic"               # in a converted comic's keywords
CODE_MARK = "BookTalker code"                 # in a converted source file's keywords
CODE_PAGE = pymupdf.Rect(0, 0, 612, 792)      # wider pages, so lines of code rarely wrap
CONVERT_VERSION = 4            # 2: RAR 4 long distances + filters fixed, Gutenberg licence dropped; 3: code kept as code;
                               # 4: exact words kept for shaped scripts (layout.EXACT_WORDS)
NEW_PAGE = "<!--bt:newpage-->"               # in our HTML: start a new page here (CSS page breaks hang the typesetter)
NO_WINDOW = 0x08000000                         # CREATE_NO_WINDOW


def supported(path):
    return os.path.splitext(path)[1].lower() in SUPPORTED


def cached(path):
    """The PDF showing this file if it needs no (more) converting, else None."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return path
    out = os.path.join(os.path.dirname(cache_dir()), "converted", f"{doc_id(path)}-v{CONVERT_VERSION}.pdf")
    if os.path.exists(out) and os.path.getsize(out) > 0:
        if ext not in (".cbr", ".cbz", ".cbt") or _marked_comic(out):
            return out                          # (comics converted before panel reading are redone)
    return None


def to_pdf_apart(path, on_page=None):
    """to_pdf in a process of its own. MEASURED: converting in the window's process froze the
    window for seconds at a time (MuPDF holds Python's lock while it typesets) - a long text
    showed a blank page and a dead Play button. on_page(n) as pages are made."""
    import multiprocessing
    import queue
    ctx = multiprocessing.get_context("spawn")
    q = ctx.Queue()
    proc = ctx.Process(target=_apart, args=(path, q), daemon=True)
    proc.start()
    try:
        while True:
            try:
                kind, value = q.get(timeout=1)
            except queue.Empty:
                if not proc.is_alive():
                    raise RuntimeError("the converter stopped unexpectedly")
                continue
            if kind == "page":
                if on_page:
                    on_page(value)
            elif kind == "done":
                return value
            else:
                raise RuntimeError(value)
    finally:
        proc.join(5)


def _apart(path, q):
    global _on_page
    _on_page = lambda n: q.put(("page", n)) if n % 5 == 0 else None
    try:
        q.put(("done", to_pdf(path)))
    except Exception as e:      # noqa: BLE001 — reported to the window
        q.put(("error", str(e) or e.__class__.__name__))


def to_pdf(path, progress=None):
    """Return the path of a PDF showing this file (converted once, then cached)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return path
    out_dir = os.path.join(os.path.dirname(cache_dir()), "converted")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, f"{doc_id(path)}-v{CONVERT_VERSION}.pdf")
    if cached(path):
        return out
    for old in (os.path.join(out_dir, doc_id(path) + ".pdf"),) + tuple(
            os.path.join(out_dir, f"{doc_id(path)}-v{v}.pdf") for v in range(2, CONVERT_VERSION)):
        try:                                    # a conversion made by an older BookTalker
            os.remove(old)
        except OSError:
            pass
    if progress:
        progress(f"Preparing {os.path.basename(path)}…")
    tmp = out + ".part.pdf"
    if path.lower().endswith(".fb2.zip"):
        ext = ".fbz"
    if ext in NATIVE:
        try:
            _native(path, tmp)
        except Exception:
            if ext != ".mobi":
                raise
            _kindle(path, tmp)                     # a KF8-only .mobi: PyMuPDF can't, our reader can
    elif ext in KINDLE:
        _kindle(path, tmp)
    elif ext in WORD:
        if not _office("word", path, tmp):
            _word_ourselves(path, tmp, ext)
    elif ext in SLIDES:
        if not _office("powerpoint", path, tmp):
            _slides_ourselves(path, tmp, ext)
    elif ext in WEB:
        _typeset(strip_gutenberg(_read_text(path), html=True), tmp, is_html=True)
    elif ext == ".tex":
        from . import readers
        _typeset(readers.tex(_read_text(path)), tmp)
    elif ext in CODE or os.path.basename(path).lower() in ("makefile", "dockerfile"):
        _code_file(path, tmp)
    elif ext in TEXT:
        _typeset(_text_to_html(strip_gutenberg(_read_text(path)), markdown=ext in (".md", ".markdown")), tmp, is_html=True)
    elif ext in OTHER:
        _other(path, tmp, ext)
    else:
        raise ValueError(f"BookTalker can't open {ext} files yet.")
    for attempt in range(6):                    # antivirus may hold a file it just saw written
        try:
            os.replace(tmp, out)
            break
        except PermissionError:
            if attempt == 5:
                raise
            import time
            time.sleep(0.3 * (attempt + 1))
    return out


# ---- formats PyMuPDF understands ---------------------------------------------------
# scripts whose letters change shape when typeset (joined, stacked, reordered): the PDF writer can
# lose them (layout.EXACT_WORDS), so their pages' exact words are read while MuPDF still has them
SHAPED = re.compile("[\u0530-\u08ff\u0900-\u0dff\u0e00-\u0fff\u1000-\u109f\u1780-\u18af"
                    "\ua980-\ua9df\ufb13-\ufdff\ufe70-\ufeff]")


def _shaped_words(tp):
    """MuPDF's own words on a text page, when it holds a shaped script (else None)."""
    return stext_words(tp) if SHAPED.search(tp.extractText()) else None


def _keep_exact(doc, exact):
    """Store, inside the PDF, the exact words of the pages whose written text lost letters."""
    keep = {}
    for i, words in exact.items():
        if not words or i >= doc.page_count:
            continue
        written = "".join(w[4] for w in doc[i].get_text("words", sort=False))
        if "".join(w[4] for w in words) != written:
            keep[str(i)] = [[round(v, 2) for v in w[:4]] + list(w[4:]) for w in words]
    if keep:
        import json
        doc.embfile_add(EXACT_WORDS, json.dumps(keep, ensure_ascii=False).encode("utf-8"),
                        desc="BookTalker: the exact words of pages whose PDF text lost letters")


def _native(path, out):
    src = pymupdf.open(path)
    exact = {}
    if src.is_reflowable:
        src.layout(rect=BOOK_PAGE, fontsize=11)
        for i in range(src.page_count):
            exact[i] = _shaped_words(src[i].get_textpage(flags=pymupdf.TEXTFLAGS_WORDS))
    pdf = pymupdf.open("pdf", src.convert_to_pdf())
    _keep_exact(pdf, exact)
    keep = ("title", "author", "subject", "keywords")
    pdf.set_metadata({k: v for k, v in (src.metadata or {}).items() if k in keep and v})
    try:                                            # keep the book's chapters (EPUB/MOBI/FB2 contents)
        toc = src.get_toc(simple=True)
        if toc:
            pdf.set_toc(toc)
    except Exception:
        pass
    pdf.save(out, garbage=3, deflate=True)
    pdf.close()
    src.close()


# ---- Microsoft Office (exact layout) -------------------------------------------------
# Only the Office process this script starts is ever closed; the user's own windows are
# never touched (PIDs that existed before are left alone).
_PS_WORD = r"""
$ErrorActionPreference = 'Stop'
$old = @(Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object Id)
$w = New-Object -ComObject Word.Application
$mine = @(Get-Process WINWORD -ErrorAction SilentlyContinue | Where-Object { $old -notcontains $_.Id } | ForEach-Object Id)
$w.Visible = $false
$w.DisplayAlerts = 0
try {
  $d = $w.Documents.Open($args[0], $false, $true, $false)
  # 17 = PDF; the last argument (1) turns the document's headings into PDF bookmarks = chapters
  $d.ExportAsFixedFormat($args[1], 17, $false, 0, 0, 1, 1, 0, $true, $true, 1)
  $d.Close(0)
} finally {
  if ($mine.Count -gt 0) {
    try { $w.Quit(0) } catch {}
    [void][Runtime.InteropServices.Marshal]::ReleaseComObject($w)
    Start-Sleep -Milliseconds 1500
    foreach ($id in $mine) { Stop-Process -Id $id -Force -ErrorAction SilentlyContinue }
  }
}
"""
_PS_POWERPOINT = r"""
$ErrorActionPreference = 'Stop'
$old = @(Get-Process POWERPNT -ErrorAction SilentlyContinue | ForEach-Object Id)
$p = New-Object -ComObject PowerPoint.Application
$mine = @(Get-Process POWERPNT -ErrorAction SilentlyContinue | Where-Object { $old -notcontains $_.Id } | ForEach-Object Id)
try {
  $pr = $p.Presentations.Open($args[0], -1, 0, 0)
  $pr.SaveAs($args[1], 32)
  $pr.Close()
} finally {
  # PowerPoint is one shared process: only shut it if we started it
  if ($mine.Count -gt 0) {
    try { $p.Quit() } catch {}
    [void][Runtime.InteropServices.Marshal]::ReleaseComObject($p)
    Start-Sleep -Milliseconds 1500
    foreach ($id in $mine) { Stop-Process -Id $id -Force -ErrorAction SilentlyContinue }
  }
}
"""


def _office(app, path, out):
    script = _PS_WORD if app == "word" else _PS_POWERPOINT
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8") as f:
        f.write(script)
        ps1 = f.name
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                            "-File", ps1, os.path.abspath(path), os.path.abspath(out)],
                           capture_output=True, timeout=300, creationflags=NO_WINDOW)
        return r.returncode == 0 and os.path.exists(out) and os.path.getsize(out) > 0
    except (OSError, subprocess.TimeoutExpired):
        return False
    finally:
        os.unlink(ps1)


# ---- our own readers (no Office needed) ------------------------------------------------
def _word_ourselves(path, out, ext):
    from . import readers
    if ext in (".docx", ".docm", ".dotx"):
        import mammoth
        with open(path, "rb") as f:
            body = mammoth.convert_to_html(f).value
    elif ext in (".doc", ".dot"):
        body = readers.doc(path)
    elif ext == ".rtf":
        body = readers.rtf(path)
    elif ext in (".odt", ".ott"):
        body = readers.odf(path)
    else:
        body = readers.wpd(path)
    _typeset(body, out)


def _slides_ourselves(path, out, ext):
    from . import readers
    slide_page = pymupdf.Rect(0, 0, 720, 405)
    if ext in (".ppt", ".pps"):
        body = readers.ppt(path)
    elif ext in (".odp", ".otp"):
        body = readers.odf(path)
    else:
        from pptx import Presentation
        parts = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = "".join(r.text for r in para.runs).strip()
                        if t:
                            texts.append(html.escape(t))
            title = texts[0] if texts else f"Slide {i}"
            parts.append(f"<h2>{title}</h2>" + "".join(f"<p>{t}</p>" for t in texts[1:]))
        body = "\n".join(parts)
    # one slide per page
    body = body.replace("<h2>", NEW_PAGE + "<h2>")
    _typeset(body, out, page=slide_page)


def _kindle(path, out):
    from . import mobi
    try:
        body, title, images = mobi.read(path)
    except mobi.ProtectedBook:
        raise ValueError("this book is copy-protected (DRM), so it can't be read aloud")
    _typeset(body, out, images=images, title=title)


def _other(path, out, ext):
    from . import readers
    if ext == ".lit":
        # Microsoft Reader books are all 'sealed' (encrypted), even free ones: opening one means
        # breaking that protection, which BookTalker doesn't do
        if b"DRMStorage" in open(path, "rb").read(400000):
            raise ValueError("Microsoft Reader (.lit) books are copy-protected (sealed), so they can't be read aloud")
        raise ValueError("this Microsoft Reader (.lit) book couldn't be opened")
    if ext in (".djvu", ".djv"):
        from . import djvu
        djvu.to_pdf(path, out)
        return
    if ext in (".cbr", ".cbz", ".cbt"):
        _comic(path, out)
        return
    if ext == ".chm":
        from . import chm
        body, title, images = chm.read(path)
        _typeset(body, out, images=images, title=title)
    elif ext in (".ods", ".ots"):
        _typeset(readers.odf(path), out)
    elif ext in (".eml", ".mht", ".mhtml"):
        _typeset(readers.eml(path), out)
    elif ext == ".txtz":
        name, data = readers.unzip_main(path, (".txt", ".md", ".html", ".htm"))
        text = data.decode("utf-8", "replace")
        _typeset(text if name.lower().endswith((".html", ".htm")) else _text_to_html(text), out)
    elif ext == ".htmlz":
        _name, data = readers.unzip_main(path, (".html", ".htm", ".xhtml"))
        _typeset(data.decode("utf-8", "replace"), out)
    elif ext == ".fbz":
        _name, data = readers.unzip_main(path, (".fb2",))
        tmp = out + ".fb2"
        open(tmp, "wb").write(data)
        try:
            _native(tmp, out)
        finally:
            os.remove(tmp)


def _marked_comic(pdf_path):
    try:
        with pymupdf.open(pdf_path) as d:
            return COMIC_MARK in ((d.metadata or {}).get("keywords") or "")
    except Exception:
        return False


def _comic(path, out):
    """Comic-book archive (.cbr = RAR, .cbz = zip, .cbt = tar; some .cbr are really zips) -> one page per picture."""
    import re
    head = open(path, "rb").read(8)
    if head.startswith(b"Rar!"):
        from . import rar
        entries = rar.files(path)
    elif head.startswith(b"PK"):
        import zipfile
        with zipfile.ZipFile(path) as z:
            entries = [(n, z.read(n)) for n in z.namelist() if not n.endswith("/")]
    else:
        import tarfile
        with tarfile.open(path) as t:
            entries = [(m.name, t.extractfile(m).read()) for m in t.getmembers() if m.isfile()]
    pages = [(n, b) for n, b in entries if n.lower().endswith(PICTURES)]
    if not pages:
        raise ValueError("no pictures could be read from this comic archive")
    natural = lambda n: [int(x) if x.isdigit() else x.lower() for x in re.split(r"(\d+)", n[0])]
    pages.sort(key=natural)
    pdf = pymupdf.open()
    for _name, data in pages:
        try:
            img = pymupdf.open(stream=data)
            pdf.insert_pdf(pymupdf.open("pdf", img.convert_to_pdf()))
            img.close()
        except Exception:
            continue                                     # a damaged picture: skip that page
    if not pdf.page_count:
        raise ValueError("no pictures could be read from this comic archive")
    # the keyword tells the page reader to follow panels and balloons (bt/comic.py)
    pdf.set_metadata({"title": os.path.splitext(os.path.basename(path))[0], "keywords": COMIC_MARK})
    pdf.save(out, garbage=3, deflate=True)
    pdf.close()


# ---- plain text / html -> typeset pages ---------------------------------------------------
GUTENBERG_START = re.compile(r"\*{3}\s*START OF (THE|THIS) PROJECT GUTENBERG[^\n]*", re.I)
GUTENBERG_END = re.compile(r"\*{3}\s*END OF (THE|THIS) PROJECT GUTENBERG", re.I)


def strip_gutenberg(text, html=False):
    """Project Gutenberg files wrap the book in an English licence (header and footer): keep
    only what lies between its '*** START OF' and '*** END OF' lines. In HTML the cut runs to
    the end of the element holding the marker, and the page head is kept."""
    a, b = GUTENBERG_START.search(text), GUTENBERG_END.search(text)
    if not a:
        return text
    start, end = a.end(), (b.start() if b and b.start() > a.end() else len(text))
    if not html:
        return text[start:end]
    close = text.find(">", text.find("</", start))          # end of the marker's element
    body = re.search(r"<body[^>]*>", text, re.I)
    head = text[:body.end()] if body and body.end() < a.start() else ""
    if b and b.start() > a.end():
        end = text.rfind("<", 0, b.start())                 # the tag that holds the end marker
    return head + text[close + 1:end] + "</body></html>"


def _code_file(path, out):
    """A source file -> pages of code in a typewriter font, marked so that every line is read as
    code (explained or read out exactly, bt/code.py)."""
    text = _read_text(path).replace("\r\n", "\n").replace("\t", "    ")
    _typeset('<pre class="src">' + html.escape(text) + "</pre>", out, page=CODE_PAGE,
             title=os.path.basename(path), keywords=CODE_MARK)


def _read_text(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "utf-16"):
        try:
            if enc == "utf-16" and not raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(raw).best()
        if best is not None:
            return str(best)
    except ImportError:
        pass
    return raw.decode("cp1252", errors="replace")


def _text_to_html(text, markdown=False):
    from .code import block_score
    text = text.replace("\r\n", "\n")
    fenced = []

    def fence(m):                               # Markdown ``` code ``` (blank lines inside kept)
        fenced.append(m.group(1))
        return f"\n\n\x00{len(fenced) - 1}\x00\n\n"
    if markdown:
        text = re.sub(r"(?ms)^[ \t]*(?:```|~~~)[^\n]*\n(.*?)^[ \t]*(?:```|~~~)[ \t]*$", fence, text)
    out = []
    for block in re.split(r"\n\s*\n", text):
        m = re.fullmatch(r"\s*\x00(\d+)\x00\s*", block)
        if m:
            out.append("<pre>" + html.escape(fenced[int(m.group(1))].rstrip("\n")) + "</pre>")
            continue
        lines = [l for l in block.split("\n") if l.strip()]
        if len(lines) >= 2 and block_score(lines) >= (0.5 if all(l[:1] in " \t" for l in lines) else 0.75):
            out.append("<pre>" + html.escape(block.strip("\n").replace("\t", "    ")) + "</pre>")   # code: lines kept
            continue
        block = block.strip()
        if not block:
            continue
        # headings: '# x' (Markdown), '= x' (AsciiDoc), '* x' (Org), a title underlined with
        # ==== or ---- (reStructuredText), or a short line that names a chapter / is in capitals
        m = re.match(r"^(#{1,6}|={1,6}|\*{1,3})\s+(\S.*)$", block)
        under = re.match(r"^([^\n]{1,80})\n[=\-~^]{3,}\s*$", block)
        if m and (markdown or m.group(1)[0] in "=*") and "\n" not in block:
            lvl = min(len(m.group(1)) + (1 if m.group(1)[0] == "#" else 0), 3)
            out.append(f"<h{lvl}>{html.escape(m.group(2))}</h{lvl}>")
            continue
        if under:
            out.append(f"<h2>{html.escape(under.group(1).strip())}</h2>")
            continue
        if "\n" not in block:
            from .readers import looks_like_heading
            if looks_like_heading(block):
                out.append(f"<h2>{html.escape(block)}</h2>")
                continue
        if markdown:
            block = re.sub(r"[*_`]{1,3}([^*_`]+)[*_`]{1,3}", r"\1", block)
        # wrapped lines inside a paragraph are joined; the reflow step does the same for PDFs
        out.append("<p>" + html.escape(" ".join(l.strip() for l in block.split("\n"))) + "</p>")
    return "\n".join(out)


CSS = """
body { font-family: serif; font-size: 11pt; line-height: 1.45; }
h1, h2, h3, h4 { font-family: sans-serif; line-height: 1.2; margin: 14pt 0 6pt 0; }
h1 { font-size: 18pt; } h2 { font-size: 15pt; } h3 { font-size: 13pt; }
p { margin: 0 0 7pt 0; text-align: justify; }
pre, code, tt, kbd, samp { font-family: monospace; }
pre { font-size: 8.5pt; line-height: 1.25; white-space: pre-wrap; text-align: left; margin: 4pt 0 9pt 0; }
pre.src { font-size: 7.5pt; margin: 0; }
"""


_on_page = None                 # to_pdf(on_page=...): told each new page while a book is typeset
VOID_TAGS = {"br", "img", "hr", "meta", "link", "input", "wbr", "col", "area", "base", "source", "embed"}


def _pieces(body, size=20000):
    """HTML -> pieces of about `size` characters, cut only between top-level elements (never
    inside a list, table or paragraph). HTML that never comes back to the top level stays whole."""
    if len(body) <= 2 * size:
        return [body]
    out, depth, start = [], 0, 0
    for m in re.finditer(r"<(/?)([A-Za-z][\w:-]*)[^>]*?(/?)>", body):
        close, tag, selfclose = m.group(1), m.group(2).lower(), m.group(3)
        if tag in VOID_TAGS or selfclose:
            continue
        depth += -1 if close else 1
        if depth < 0:                       # stray closing tag: give up cutting, keep it whole
            return [body]
        if depth == 0 and close and m.end() - start >= size:
            out.append(body[start:m.end()])
            start = m.end()
    out.append(body[start:])
    # (nothing visible in a piece: no page for it)
    return [p for p in out if re.sub(r"<[^>]+>|\s|&nbsp;", "", p) or "<img" in p]


def _draw_reading(story, dev, rect):
    """Draw a placed story onto a PDF page, reading MuPDF's own letters on the way: the page is
    recorded once, then played to the PDF and to a text page. -> _shaped_words of that page."""
    from pymupdf import mupdf
    box = mupdf.FzRect(*rect)
    everything = mupdf.FzRect(mupdf.FzRect.Fixed_INFINITE)
    recording = mupdf.FzDisplayList(box)
    rec = pymupdf.DeviceWrapper(recording)
    story.draw(rec)
    mupdf.fz_close_device(rec.this)
    mupdf.fz_run_display_list(recording, dev.this, mupdf.FzMatrix(), everything, mupdf.FzCookie())
    text = mupdf.FzStextPage(box)
    reader = pymupdf.DeviceWrapper(text, pymupdf.TEXTFLAGS_WORDS)
    mupdf.fz_run_display_list(recording, reader.this, mupdf.FzMatrix(), everything, mupdf.FzCookie())
    mupdf.fz_close_device(reader.this)
    return _shaped_words(pymupdf.TextPage(text))


def _typeset(body_html, out, is_html=True, page=BOOK_PAGE, images=None, title=None, keywords=None):
    """HTML -> book pages. Headings (h1-h3) become the PDF's chapters; images in `images`
    (name -> bytes) are shown, scaled to the page width at most."""
    body_html = re.sub(r"(?is)<(script|style|nav|header|footer)\b.*?</\1>", "", body_html)
    # PyMuPDF's typesetter never finishes on CSS page breaks (measured): drop them
    body_html = re.sub(r"(?i)page-break-(before|after|inside)\s*:\s*[\w-]+\s*;?", "", body_html)
    from .readers import promote_headings
    body_html = promote_headings(body_html)
    margin = 54 if page.width < 600 else 40
    where = page + (margin, margin, -margin, -margin)
    archive = None
    if images:
        archive = pymupdf.Archive()
        from PIL import Image
        import io
        sizes = {}
        for name, data in images.items():
            archive.add(data, name)
            try:
                w, _h = Image.open(io.BytesIO(data)).size
                sizes[name] = min(w * 0.75, where.width)
            except Exception:
                pass
        body_html = re.sub(r'<img\b([^>]*?)src="([^"]+)"([^>]*)>',
                           lambda m: (f'<img src="{m.group(2)}" width="{int(sizes[m.group(2)])}">'
                                      if m.group(2) in sizes else ""), body_html)
    else:
        body_html = re.sub(r"<img\b[^>]*>", "", body_html)
    import io
    buf = io.BytesIO()                                  # typeset in memory, write the file once
    writer = pymupdf.DocumentWriter(buf)
    heads = []
    exact = {}

    def record(pos):
        if pos.heading and (pos.open_close & 1) and pos.text and pos.text.strip():
            heads.append((pos.heading, " ".join(pos.text.split()), pos.page_num, pos.rect[1]))
    n = 0
    limit = 50 + len(body_html) // 40                   # far more pages than the text could need
    dev = None
    for section in body_html.split(NEW_PAGE):           # each section starts on a new page
        if not section.strip():
            continue
        # MEASURED: each place() works through the whole rest of its story (a 236-page book: 42 ms
        # a page, 883 pages: 175 ms), so a long text goes in as pieces, each carrying on where
        # the last one stopped on the page
        top = where.y0
        pieces = _pieces(section)
        for i, piece in enumerate(pieces):
            # the page's body margin belongs to the whole text, not to each piece (MEASURED: with
            # it, every seam opened a 15 pt gap; without it the lines sit exactly as in one piece)
            css = CSS + ("" if i == 0 else " body { margin-top: 0 }") + ("" if i == len(pieces) - 1 else " body { margin-bottom: 0 }")
            story = pymupdf.Story(html=piece, user_css=css, archive=archive)
            more, stuck = True, 0
            while more and n < limit:
                if dev is None:
                    n += 1
                    dev, top = writer.begin_page(page), where.y0
                    if _on_page:
                        _on_page(n)
                more, filled = story.place(pymupdf.Rect(where.x0, top, where.x1, where.y1))
                story.element_positions(record, {"page_num": n})
                words = _draw_reading(story, dev, page)
                if words:                               # (pieces sharing a page: blocks numbered on)
                    have = exact.get(n - 1) or []
                    shift = 1 + max((w[5] for w in have), default=-1)
                    exact[n - 1] = have + [w[:5] + (w[5] + shift,) + tuple(w[6:]) for w in words]
                filled = pymupdf.Rect(filled)
                # (MEASURED: a story started in a strip shorter than a few lines lost whole
                # paragraphs - MuPDF counts them placed but draws nothing - so a piece that ends
                # near the foot of the page closes it)
                if more or filled.is_empty or filled.y1 >= where.y1 - 60:
                    writer.end_page()
                    dev = None
                else:
                    top = filled.y1                     # the next piece starts below this one
                stuck = stuck + 1 if filled.is_empty else 0
                if stuck >= 3:                          # something that can never fit: move on
                    break
        if dev is not None:
            writer.end_page()
            dev = None
    if n == 0:
        writer.begin_page(page)
        writer.end_page()
    writer.close()
    doc = pymupdf.open("pdf", buf.getvalue())
    from .book import CHAPTER_WORD, DIVISION
    # the typesetter drops the space where a long heading wraps: take the titles from the HTML
    src = [" ".join(html.unescape(re.sub(r"<[^>]+>", " ", t)).split())
           for t in re.findall(r"(?is)<h[1-3]\b[^>]*>(.*?)</h[1-3]>", body_html)]
    src = [t for t in src if t]
    if len(src) == len(heads):
        heads = [(lvl, t, pg, y) for (lvl, _old, pg, y), t in zip(heads, src)]
    if sum(1 for h in heads if DIVISION.search(h[1])) >= 2:    # a chaptered book: chapters (and
        heads = [h for h in heads if CHAPTER_WORD.search(h[1])]  # foreword etc.), not title-page lines
    toc, last = [], 0
    top = min((h[0] for h in heads), default=1)
    for lvl, text, pg, y in heads:                      # the headings become the book's chapters
        lvl = max(1, min(lvl - top + 1, last + 1))
        toc.append([lvl, text[:120], pg, {"kind": pymupdf.LINK_GOTO, "page": pg - 1, "to": pymupdf.Point(0, y)}])
        last = lvl
    try:
        if toc:
            doc.set_toc(toc)
    except Exception:
        pass
    if title or keywords:
        doc.set_metadata({k: v for k, v in (("title", title), ("keywords", keywords)) if v})
    _keep_exact(doc, exact)
    doc.save(out, garbage=3, deflate=True)
    doc.close()
